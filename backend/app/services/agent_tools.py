"""Инструменты агента для работы с закупками."""
import uuid
import hashlib
import zipfile
import json
import io
from datetime import datetime
from typing import Dict, Any, List
from sqlalchemy.orm import Session
from fastapi import HTTPException
from docx import Document as DocxDocument
from app.models.case import Case
from app.models.document import Document
from app.models.document_category import DocumentCategory
from app.models.template import Template
from app.models.supplier import Supplier
from app.models.procurement import (
    CaseCandidate,
    CommercialOffer,
    SupplierSecurityReview
)
from app.services.storage import get_s3_client, ensure_bucket_exists, upload_file, get_file
from app.core.settings import settings


def validate_procurement(db: Session, case_id: uuid.UUID) -> Dict[str, Any]:
    """
    Валидация закупки: проверка наличия необходимых данных.
    
    Возвращает:
    {
        "ok": bool,
        "missing": List[str],  # Что отсутствует
        "blockers": List[str]  # Критические проблемы, блокирующие процесс
    }
    """
    case = db.query(Case).filter(Case.id == case_id).first()
    if not case:
        return {
            "ok": False,
            "missing": [f"Кейс с ID {case_id} не найден"],
            "blockers": [f"Кейс с ID {case_id} не найден"]
        }
    
    missing = []
    blockers = []
    
    # Базовые данные кейса
    if not case.subject:
        missing.append("Предмет закупки (subject)")
        blockers.append("Не указан предмет закупки")
    
    if not case.initiator_department:
        missing.append("Инициатор (initiator_department)")
    
    if not case.planned_deadline:
        missing.append("Плановый срок (planned_deadline)")
    
    # Кандидаты
    candidates_count = db.query(CaseCandidate).filter(
        CaseCandidate.case_id == case_id
    ).count()
    
    if candidates_count == 0:
        missing.append("Кандидаты (candidates)")
        blockers.append("Нет кандидатов для закупки")
    
    # Коммерческие предложения
    offers_count = db.query(CommercialOffer).filter(
        CommercialOffer.case_id == case_id
    ).count()
    
    if offers_count == 0:
        missing.append("Коммерческие предложения (commercial_offers)")
    elif offers_count < 3:
        missing.append(f"Недостаточно КП: {offers_count} из 3 минимум")
    
    # Валидация НМЦ
    if not case.is_validation_ok:
        if case.offers_count < 3:
            blockers.append(f"Недостаточно КП для валидации НМЦ: {case.offers_count} из 3")
        elif case.validation_coeff and case.validation_coeff > 0.33:
            blockers.append(f"Коэффициент валидации слишком высокий: {case.validation_coeff} (максимум 0.33)")
        else:
            missing.append("Валидация НМЦ не пройдена")
    
    # Проверка СБ (если выбран поставщик)
    if case.selected_supplier_id:
        security_review = db.query(SupplierSecurityReview).filter(
            SupplierSecurityReview.case_id == case_id,
            SupplierSecurityReview.supplier_id == case.selected_supplier_id
        ).first()
        
        if not security_review:
            blockers.append("Нет проверки СБ для выбранного поставщика")
        elif security_review.status != "APPROVED":
            blockers.append(f"Проверка СБ не пройдена: статус {security_review.status}")
    
    ok = len(blockers) == 0
    
    return {
        "ok": ok,
        "missing": missing,
        "blockers": blockers
    }


def export_zip(db: Session, case_id: uuid.UUID, mode: str = "standard") -> Dict[str, Any]:
    """
    Экспорт кейса в ZIP архив.
    
    Args:
        db: Database session
        case_id: ID кейса
        mode: Режим экспорта ("standard", "full")
    
    Returns:
        {"ok": bool, "document_id": str | None, "reason": str | None}
    """
    case = db.query(Case).filter(Case.id == case_id).first()
    if not case:
        return {
            "ok": False,
            "document_id": None,
            "reason": f"Кейс с ID {case_id} не найден"
        }
    
    # Получаем категорию EXPORT_ZIP (создаём если нет)
    export_category = db.query(DocumentCategory).filter(
        DocumentCategory.key == "EXPORT_ZIP"
    ).first()
    if not export_category:
        from app.services.document_category import create_category
        export_category = create_category(db, "Экспорт ZIP", "EXPORT_ZIP")
    
    # Получаем все документы кейса, кроме EXPORT_ZIP (не вкладываем zip в zip)
    documents = db.query(Document).filter(
        Document.case_id == case_id,
        Document.category_id != export_category.id
    ).all()
    
    if not documents:
        return {
            "ok": False,
            "document_id": None,
            "reason": "Нет документов для экспорта"
        }
    
    # Создаём ZIP в памяти
    zip_buffer = io.BytesIO()
    manifest_files = []
    
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for doc in documents:
            try:
                # Скачиваем файл из MinIO
                file_content = get_file(doc.storage_bucket, doc.storage_key)
                
                # Вычисляем SHA256
                sha256_hash = hashlib.sha256(file_content).hexdigest()
                
                # Добавляем в ZIP
                zip_file.writestr(doc.original_filename, file_content)
                
                # Добавляем в manifest
                manifest_files.append({
                    "document_id": str(doc.id),
                    "filename": doc.original_filename,
                    "category_key": doc.category.key if doc.category else None,
                    "storage_key": doc.storage_key,
                    "sha256": sha256_hash,
                    "size_bytes": doc.size_bytes or len(file_content)
                })
            except Exception as e:
                # Пропускаем файлы, которые не удалось скачать
                continue
        
        # Создаём manifest.json
        manifest = {
            "case_id": str(case_id),
            "case_code": case.code,
            "exported_at": datetime.utcnow().isoformat(),
            "mode": mode,
            "files": manifest_files
        }
        manifest_json = json.dumps(manifest, ensure_ascii=False, indent=2)
        zip_file.writestr("manifest.json", manifest_json.encode('utf-8'))
    
    zip_buffer.seek(0)
    zip_content = zip_buffer.read()
    zip_size = len(zip_content)
    
    # Вычисляем SHA256 для ZIP
    zip_sha256 = hashlib.sha256(zip_content).hexdigest()
    
    # Сохраняем ZIP в MinIO
    ensure_bucket_exists(settings.minio_bucket_files)
    zip_filename = f"{case.code}_export_{mode}.zip"
    storage_key = f"cases/{case.code}/EXPORT/{uuid.uuid4().hex}_{zip_filename}"
    
    upload_file(
        settings.minio_bucket_files,
        storage_key,
        zip_content,
        "application/zip"
    )
    
    # Создаём Document для ZIP
    zip_document = Document(
        case_id=case_id,
        category_id=export_category.id,
        doc_type="CASE",
        original_filename=zip_filename,
        content_type="application/zip",
        size_bytes=zip_size,
        storage_bucket=settings.minio_bucket_files,
        storage_key=storage_key,
        status="DONE"
    )
    db.add(zip_document)
    db.commit()
    db.refresh(zip_document)
    
    return {
        "ok": True,
        "document_id": str(zip_document.id),
        "reason": None
    }


def generate_doc(db: Session, case_id: uuid.UUID, template_key: str) -> Dict[str, Any]:
    """
    Генерация документа по шаблону.
    
    Args:
        db: Database session
        case_id: ID кейса
        template_key: Ключ шаблона (например, "NMC_REFERENCE")
    
    Returns:
        {"ok": bool, "document_id": str | None, "reason": str | None}
    """
    case = db.query(Case).filter(Case.id == case_id).first()
    if not case:
        return {
            "ok": False,
            "document_id": None,
            "reason": f"Кейс с ID {case_id} не найден"
        }
    
    # Получаем шаблон из БД
    template = db.query(Template).filter(Template.key == template_key).first()
    if not template or not template.storage_key:
        return {
            "ok": False,
            "document_id": None,
            "reason": f"Шаблон с ключом '{template_key}' не найден"
        }
    
    # Получаем категорию GENERATED_DOC (создаём если нет)
    generated_category = db.query(DocumentCategory).filter(
        DocumentCategory.key == "GENERATED_DOC"
    ).first()
    if not generated_category:
        from app.services.document_category import create_category
        generated_category = create_category(db, "Сгенерированный документ", "GENERATED_DOC")
    
    # Скачиваем шаблон из MinIO
    try:
        template_content = get_file(settings.minio_bucket_templates, template.storage_key)
    except Exception as e:
        return {
            "ok": False,
            "document_id": None,
            "reason": f"Ошибка загрузки шаблона: {str(e)}"
        }
    
    # Открываем docx
    try:
        doc = DocxDocument(io.BytesIO(template_content))
    except Exception as e:
        return {
            "ok": False,
            "document_id": None,
            "reason": f"Ошибка чтения шаблона: {str(e)}"
        }
    
    # Получаем выбранного поставщика если есть
    selected_supplier_name = None
    if case.selected_supplier_id:
        supplier = db.query(Supplier).filter(Supplier.id == case.selected_supplier_id).first()
        if supplier:
            selected_supplier_name = supplier.name
    
    # Заменяем плейсхолдеры в документе
    replacements = {
        "[[CASE_CODE]]": case.code or "",
        "[[CASE_TITLE]]": case.title or "",
        "[[NMC_AVG_PRICE]]": str(float(case.nmc_avg_price)) if case.nmc_avg_price else "0.00",
        "[[VALIDATION_COEFF]]": str(float(case.validation_coeff)) if case.validation_coeff else "0.00",
        "[[OFFERS_COUNT]]": str(case.offers_count) if case.offers_count else "0",
        "[[SELECTED_SUPPLIER_NAME]]": selected_supplier_name or "не выбран"
    }
    
    # Заменяем в параграфах
    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)
    
    # Заменяем в таблицах
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old, new in replacements.items():
                        if old in paragraph.text:
                            paragraph.text = paragraph.text.replace(old, new)
    
    # Сохраняем результат в BytesIO
    output_buffer = io.BytesIO()
    doc.save(output_buffer)
    output_buffer.seek(0)
    generated_content = output_buffer.read()
    
    # Сохраняем в MinIO
    ensure_bucket_exists(settings.minio_bucket_files)
    output_filename = f"{template_key}_{case.code}.docx"
    storage_key = f"cases/{case.code}/GENERATED/{uuid.uuid4().hex}_{output_filename}"
    
    upload_file(
        settings.minio_bucket_files,
        storage_key,
        generated_content,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    
    # Создаём Document
    generated_document = Document(
        case_id=case_id,
        category_id=generated_category.id,
        doc_type="CASE",
        original_filename=output_filename,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=len(generated_content),
        storage_bucket=settings.minio_bucket_files,
        storage_key=storage_key,
        status="DONE"
    )
    db.add(generated_document)
    db.commit()
    db.refresh(generated_document)
    
    return {
        "ok": True,
        "document_id": str(generated_document.id),
        "reason": None
    }

