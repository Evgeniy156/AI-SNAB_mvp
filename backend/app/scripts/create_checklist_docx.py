"""Временный скрипт для создания валидного checklist_procurement.docx."""
from pathlib import Path
from docx import Document

def create_checklist_docx():
    """Создаёт минимальный валидный .docx файл для чек-листа."""
    app_dir = Path(__file__).parent.parent
    docx_path = app_dir / "docs" / "checklist_procurement.docx"
    
    # Создаём документ
    doc = Document()
    
    # Заголовок
    doc.add_heading('Чек-лист закупочной процедуры', 0)
    
    # Секция I
    doc.add_heading('I. Подготовительный этап', 1)
    doc.add_paragraph('1. Основание закупки')
    doc.add_paragraph('☐ Служебная записка')
    doc.add_paragraph('☐ Техническое задание')
    doc.add_paragraph('☐ Обоснование закупки')
    
    # Секция II
    doc.add_heading('II. Анализ рынка', 1)
    doc.add_paragraph('2. Анализ рынка')
    doc.add_paragraph('☐ Изучение рынка поставщиков')
    doc.add_paragraph('☐ Сравнение предложений')
    
    # Сохраняем
    doc.save(docx_path)
    print(f"Создан файл: {docx_path}")
    print(f"Размер: {docx_path.stat().st_size} байт")

if __name__ == "__main__":
    create_checklist_docx()

